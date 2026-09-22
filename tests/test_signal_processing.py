from __future__ import annotations

import unittest
import shutil
import uuid
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PySide6.QtCore import QUrl

from signal_analysis import (
    AnalysisSettings,
    SignalAnalysisPage,
    TraceProcessingSettings,
    analyze_trace,
    analyze_trace_batch,
    detect_signal_peaks,
    estimate_baseline,
    first_detected_peak_time,
    is_auxiliary_trace_column,
    parse_trace_name,
    reference_match_key,
    validation_recording_match_key,
    prepare_signal,
    selected_result_frames,
    summary_display_columns,
    write_metrics_workbook,
    write_processed_workbook,
)


class SignalProcessingTests(unittest.TestCase):
    def test_signal_processing_uses_imported_units_without_scale_or_offset_controls(self):
        settings = AnalysisSettings()
        processing = TraceProcessingSettings()

        self.assertFalse(hasattr(settings, "scale"))
        self.assertFalse(hasattr(settings, "offset"))
        self.assertFalse(hasattr(processing, "scale"))

    def test_trace_drop_accepts_supported_files_and_rejects_other_formats(self):
        urls = [
            QUrl.fromLocalFile(str(Path("C:/traces/a.CSV"))),
            QUrl.fromLocalFile(str(Path("C:/traces/results.xlsx"))),
            QUrl.fromLocalFile(str(Path("C:/traces/readme.txt"))),
            QUrl.fromLocalFile(str(Path("C:/traces/a.CSV"))),
        ]

        selected, rejected = SignalAnalysisPage.trace_paths_from_urls(urls)

        self.assertEqual([path.name for path in selected], ["a.CSV", "results.xlsx"])
        self.assertEqual([path.name for path in rejected], ["readme.txt"])

    def test_batch_analysis_returns_reference_normalization(self):
        time = np.arange(0.0, 4.0, 0.02)
        pulse = sum(np.exp(-((time - center) / 0.09) ** 2) for center in (0.5, 1.5, 2.5, 3.5))
        selected = "Drug | B-1_Force_uN"
        reference = "Control | B-1_Force_uN"
        settings = AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)
        result = analyze_trace_batch(
            time,
            {selected: pulse, reference: pulse * 0.5},
            {selected: settings, reference: settings},
            [selected],
            [reference],
            {reference},
        )
        self.assertEqual(result["selected_count"], 1)
        self.assertEqual(result["summary"].iloc[0]["Reference_trace"], reference)
        self.assertAlmostEqual(
            result["summary"].iloc[0]["Amplitude_mean_pct_baseline"], 200.0
        )
        self.assertAlmostEqual(
            result["summary"].iloc[0]["Peak_force_mean_pct_baseline"], 200.0
        )
        self.assertAlmostEqual(
            result["summary"].iloc[0]["Max_contraction_slope_per_s_mean_pct_baseline"],
            200.0,
        )
        self.assertAlmostEqual(
            result["summary"].iloc[0]["Max_relaxation_slope_per_s_mean_pct_baseline"],
            200.0,
        )
        self.assertIn("Rise_10_90_ms_mean_pct_baseline", result["summary"].columns)
        self.assertIn("CTD50_ms_mean_pct_baseline", result["summary"].columns)

    def test_reference_matching_ignores_condition_prefix(self):
        time = np.arange(0.0, 4.0, 0.02)
        pulse = sum(
            np.exp(-((time - center) / 0.09) ** 2)
            for center in (0.5, 1.5, 2.5, 3.5)
        )
        reference = "StimTrace | BL_B-1_Force_uN"
        selected = "StimTrace | H_B-1_Force_uN"
        settings = AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)
        result = analyze_trace_batch(
            time,
            {reference: pulse * 0.5, selected: pulse},
            {reference: settings, selected: settings},
            [selected],
            [reference],
            {reference},
        )
        row = result["summary"].iloc[0]
        self.assertEqual(reference_match_key(reference), "b-1")
        self.assertEqual(reference_match_key(selected), "b-1")
        self.assertEqual(row["Reference_trace"], reference)
        self.assertEqual(row["Reference_match_status"], "Matched by specimen")
        self.assertAlmostEqual(row["Amplitude_mean_pct_baseline"], 200.0)

    def test_summary_includes_normalized_columns_when_references_are_matched(self):
        columns = summary_display_columns(pd.DataFrame({
            "Trace": ["H_B-1"],
            "Peak_force_mean": [1250.0],
            "Peak_force_mean_pct_baseline": [125.0],
            "Amplitude_mean": [100.0],
            "Amplitude_mean_pct_baseline": [125.0],
            "Beat_rate_BPM": [40.0],
            "Beat_rate_BPM_pct_baseline": [110.0],
            "Relaxation_50_ms_mean": [300.0],
            "Relaxation_50_ms_mean_pct_baseline": [105.0],
            "Rise_10_90_ms_mean": [150.0],
            "Rise_10_90_ms_mean_pct_baseline": [95.0],
            "CTD50_ms_mean": [400.0],
            "CTD50_ms_mean_pct_baseline": [115.0],
            "Max_contraction_slope_per_s_mean": [2000.0],
            "Max_contraction_slope_per_s_mean_pct_baseline": [130.0],
            "Max_relaxation_slope_per_s_mean": [-1200.0],
            "Max_relaxation_slope_per_s_mean_pct_baseline": [90.0],
        }))
        self.assertIn("Amplitude_mean_pct_baseline", columns)
        self.assertIn("Peak_force_mean_pct_baseline", columns)
        self.assertIn("Beat_rate_BPM_pct_baseline", columns)
        self.assertIn("Relaxation_50_ms_mean_pct_baseline", columns)
        self.assertIn("Rise_10_90_ms_mean_pct_baseline", columns)
        self.assertIn("CTD50_ms_mean_pct_baseline", columns)
        self.assertIn("Max_contraction_slope_per_s_mean_pct_baseline", columns)
        self.assertIn("Max_relaxation_slope_per_s_mean_pct_baseline", columns)

    def test_summary_hides_normalized_columns_without_reference_results(self):
        columns = summary_display_columns(pd.DataFrame({
            "Trace": ["H_B-1"],
            "Amplitude_mean": [100.0],
        }))
        self.assertNotIn("Amplitude_mean_pct_baseline", columns)

    def test_loaded_source_folder_label_is_clear_for_one_or_multiple_folders(self):
        one_folder = SignalAnalysisPage.loaded_source_folder_text([
            Path("C:/traces/a.csv"), Path("C:/traces/b.csv"),
        ])
        multiple_folders = SignalAnalysisPage.loaded_source_folder_text([
            Path("C:/traces/a.csv"), Path("D:/other/b.csv"),
        ])
        self.assertTrue(one_folder.endswith("traces"))
        self.assertIn("Multiple source folders", multiple_folders)

    def test_export_default_path_is_beside_loaded_trace_file(self):
        page = SignalAnalysisPage.__new__(SignalAnalysisPage)
        page.path = Path("C:/traces/force.csv")
        page.paths = [page.path]
        self.assertEqual(
            page.export_default_path("force_metrics.xlsx"),
            str(Path("C:/traces/force_metrics.xlsx")),
        )

    def test_metrics_workbook_writer_is_reopenable(self):
        folder = Path.cwd() / f".test-metrics-workbook-{uuid.uuid4().hex}"
        folder.mkdir()
        self.addCleanup(shutil.rmtree, folder, True)
        target = folder / "metrics.xlsx"
        saved = write_metrics_workbook(
            str(target),
            pd.DataFrame({"Trace": ["A"]}),
            pd.DataFrame({"Trace": ["A"], "Beat": [1]}),
            pd.DataFrame({"Reference trace": ["None selected"]}),
            pd.DataFrame({"Setting": ["smoothing_s"], "Value": [0.15]}),
            pd.DataFrame({"Trace": ["A"], "invert": [False]}),
        )
        self.assertEqual(Path(saved), target)
        with pd.ExcelFile(target) as workbook:
            self.assertEqual(
                workbook.sheet_names,
                ["Trace summary", "Beat metrics", "Reference traces", "Settings", "Trace processing"],
            )
        workbook = load_workbook(target, read_only=False)
        self.addCleanup(workbook.close)
        self.assertTrue(
            all(sheet.freeze_panes == "B2" for sheet in workbook.worksheets)
        )

    def test_single_peak_does_not_claim_a_beat_rate(self):
        time = np.arange(0.0, 2.0, 0.05)
        signal = np.exp(-((time - 1.0) / 0.12) ** 2)
        summary, beats, _ = analyze_trace(
            time,
            signal,
            "single",
            AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0),
        )
        self.assertEqual(len(beats), 1)
        self.assertTrue(np.isnan(summary["Beat_rate_BPM"]))

    def test_absolute_prominence_rejects_subthreshold_jitter(self):
        time = np.arange(0.0, 4.0, 0.01)
        jitter = 0.08 * np.sin(2 * np.pi * 8 * time)

        with self.assertRaisesRegex(ValueError, "No beats detected"):
            analyze_trace(
                time,
                jitter,
                "jitter",
                AnalysisSettings(
                    smoothing_s=0.0,
                    prominence_mode="absolute",
                    prominence_absolute=0.5,
                ),
            )

    def test_batch_summary_retains_trace_when_no_beats_are_detected(self):
        time = np.arange(0.0, 4.0, 0.01)
        trace_name = "StimTrace | CTRL_A-1_Force_uN"
        jitter = 0.08 * np.sin(2 * np.pi * 8 * time)
        settings = AnalysisSettings(
            smoothing_s=0.0,
            prominence_mode="absolute",
            prominence_absolute=0.5,
        )

        result = analyze_trace_batch(
            time,
            {trace_name: jitter},
            {trace_name: settings},
            [trace_name],
            [],
            set(),
        )

        self.assertEqual(len(result["summary"]), 1)
        row = result["summary"].iloc[0]
        self.assertEqual(row["Trace"], trace_name)
        self.assertEqual(row["N_beats"], 0)
        self.assertEqual(row["Analysis_status"], "No beats detected")
        self.assertIn("No beats detected", row["Analysis_error"])
        self.assertTrue(np.isnan(row["Peak_force_mean"]))
        self.assertTrue(result["beats"].empty)
        self.assertIn("Analysis_status", summary_display_columns(result["summary"]))

    def test_absolute_prominence_must_be_positive(self):
        time = np.arange(0.0, 2.0, 0.01)
        signal = np.sin(2 * np.pi * time)

        with self.assertRaisesRegex(ValueError, "greater than zero"):
            analyze_trace(
                time,
                signal,
                "signal",
                AnalysisSettings(
                    smoothing_s=0.0,
                    prominence_mode="absolute",
                    prominence_absolute=0.0,
                ),
            )

    def test_rolling_percentile_baseline_has_finite_edges(self):
        values = np.linspace(0.0, 1.0, 21)
        baseline = estimate_baseline(
            values,
            0.1,
            AnalysisSettings(
                baseline_method="rolling_percentile",
                baseline_window_s=0.5,
                baseline_percentile=10.0,
            ),
        )
        self.assertTrue(np.isfinite(baseline).all())

    def test_auxiliary_trace_filter_avoids_baseline_condition_names(self):
        self.assertTrue(is_auxiliary_trace_column("Trace A | baseline"))
        self.assertTrue(is_auxiliary_trace_column("Trace_A_uncorrected"))
        self.assertTrue(is_auxiliary_trace_column("Trace A | original"))
        self.assertFalse(is_auxiliary_trace_column("BaselineClean | Trace_A corrected"))
        self.assertFalse(is_auxiliary_trace_column("Treatment | Trace_A_Force_uN"))

    def test_metric_export_frames_follow_current_selection(self):
        summary = pd.DataFrame({"Trace": ["A", "B"], "N_beats": [3, 4]})
        beats = pd.DataFrame({"Trace": ["A", "A", "B"], "Peak": [1, 2, 3]})
        selected_summary, selected_beats = selected_result_frames(
            summary,
            beats,
            ["B"],
        )
        self.assertEqual(selected_summary["Trace"].tolist(), ["B"])
        self.assertEqual(selected_beats["Trace"].tolist(), ["B"])

    def test_processed_trace_workbook_can_be_reopened_without_sheet_dialog(self):
        folder = Path.cwd() / f".test-processed-workbook-{uuid.uuid4().hex}"
        folder.mkdir()
        self.addCleanup(shutil.rmtree, folder, True)
        path = folder / "selected_processed_traces.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(
                {"time_s": [0.0, 0.1, 0.2], "Trace A | corrected": [0.0, 1.0, 0.0]}
            ).to_excel(writer, sheet_name="Processed traces", index=False)
            pd.DataFrame({"Setting": ["smoothing_s"], "Value": [0.15]}).to_excel(
                writer, sheet_name="Settings", index=False
            )
        page_stub = type("PageStub", (), {"_preloaded_trace_files": {}})()
        loaded = SignalAnalysisPage._read_trace_file(
            page_stub,
            path,
            allow_dialog=False,
        )
        self.assertIsNotNone(loaded)
        data, time_column = loaded
        self.assertEqual(time_column, "time_s")
        self.assertIn("Trace A | corrected", data.columns)

    def test_processed_export_contains_only_original_and_processed_signals(self):
        folder = Path.cwd() / f".test-processed-export-{uuid.uuid4().hex}"
        folder.mkdir()
        self.addCleanup(shutil.rmtree, folder, True)
        path = folder / "processed.xlsx"
        time = np.linspace(0.0, 1.0, 11)
        signal = 5.0 + np.sin(time * np.pi)
        write_processed_workbook(
            str(path),
            "time_s",
            time,
            {"Trace A": signal},
            {
                "Trace A": AnalysisSettings(
                    smoothing_s=0.0,
                    baseline_method="constant_percentile",
                    time_shift_frames=2,
                )
            },
            pd.DataFrame(),
            pd.DataFrame(),
            pd.DataFrame([{"Trace": "Trace A", "time_shift_frames": 2}]),
        )
        exported = pd.read_excel(path, sheet_name="Processed traces")
        self.assertEqual(
            exported.columns.tolist(),
            ["time_s", "Trace A | original", "Trace A | processed"],
        )
        self.assertAlmostEqual(exported["time_s"].iloc[0], 0.2)
        processing = pd.read_excel(path, sheet_name="Trace processing")
        self.assertEqual(processing.loc[0, "time_shift_frames"], 2)
        workbook = load_workbook(path, read_only=False)
        self.addCleanup(workbook.close)
        self.assertTrue(
            all(sheet.freeze_panes == "B2" for sheet in workbook.worksheets)
        )

    def test_generated_trace_names_include_the_analysis_method(self):
        self.assertEqual(
            SignalAnalysisPage.trace_display_name(
                Path("B-2_point_tracking_force.csv"),
                "Point_1_Force_uN",
                multiple=False,
                duplicate_stems=set(),
            ),
            "Point Tracking | B-2_Point_1_Force_uN",
        )
        self.assertEqual(
            SignalAnalysisPage.trace_display_name(
                Path("stimtrace_force_traces.csv"),
                "B-2_Force_uN",
                multiple=False,
                duplicate_stems=set(),
            ),
            "StimTrace | B-2_Force_uN",
        )
        self.assertEqual(
            SignalAnalysisPage.trace_display_name(
                Path("combined_force_results.csv"),
                "B-2_Force_uN",
                multiple=True,
                duplicate_stems=set(),
            ),
            "StimTrace | B-2_Force_uN",
        )

    def test_inversion_and_baseline_settings_are_trace_specific(self):
        page_stub = type(
            "PageStub",
            (),
            {
                "trace_processing": {
                    "trace_a": TraceProcessingSettings(
                        invert=True,
                        baseline_method="rolling_percentile",
                        baseline_percentile=12.0,
                        baseline_window_s=3.0,
                        time_shift_frames=4,
                    )
                }
            },
        )()
        shared = AnalysisSettings(smoothing_s=0.25, invert=False, baseline_method="none")
        trace_a = SignalAnalysisPage.settings_for_column(page_stub, "trace_a", shared)
        trace_b = SignalAnalysisPage.settings_for_column(page_stub, "trace_b", shared)
        self.assertTrue(trace_a.invert)
        self.assertEqual(trace_a.baseline_method, "rolling_percentile")
        self.assertEqual(trace_a.baseline_window_s, 3.0)
        self.assertEqual(trace_a.time_shift_frames, 4)
        self.assertFalse(trace_b.invert)
        self.assertEqual(trace_b.baseline_method, "none")
        self.assertEqual(trace_b.smoothing_s, 0.25)

    def test_trace_name_parsing(self):
        self.assertEqual(
            parse_trace_name("Treatment | specimen_1_Force_uN"),
            ("Treatment", "specimen_1", ""),
        )

    def test_reference_matching_ignores_processed_trace_labels(self):
        self.assertEqual(
            reference_match_key("StimTrace | BL_B-1_Force_uN | processed"),
            "b-1",
        )
        self.assertEqual(
            reference_match_key("StimTrace | BL_B-1_Force_uN | smoothed"),
            "b-1",
        )
        self.assertEqual(
            reference_match_key("StimTrace | H_B-1_Force_uN | corrected"),
            "b-1",
        )

    def test_reference_matching_pairs_validation_methods_by_original_recording(self):
        manual = "1.5_Hz_Ramp_B-2_ortho_manual_ellipse_validation | manual_ellipse_force_uN"
        dl = "1.5_Hz_Ramp_B-2_ortho_manual_ellipse_validation | dl_raw_force_uN | corrected"
        points = "1.5_Hz_Ramp_B-2_ortho_point_tracking_validation | point_force_uN"
        expected = "recording::1.5_hz_ramp_b-2_ortho"
        self.assertEqual(validation_recording_match_key(manual), "1.5_hz_ramp_b-2_ortho")
        self.assertEqual(reference_match_key(manual), expected)
        self.assertEqual(reference_match_key(dl), expected)
        self.assertEqual(reference_match_key(points), expected)

        compact_manual = "Validation | 1.5_Hz_Ramp_B-2_ortho_manual_ellipse_Force_uN"
        compact_dl = "Validation | 1.5_Hz_Ramp_B-2_ortho_dl_raw_Force_uN"
        compact_point = "Validation | 1.5_Hz_Ramp_B-2_ortho_Point_1_Force_uN"
        self.assertEqual(reference_match_key(compact_manual), expected)
        self.assertEqual(reference_match_key(compact_dl), expected)
        self.assertEqual(reference_match_key(compact_point), expected)

        time = np.arange(0.0, 4.0, 0.02)
        pulse = sum(np.exp(-((time - center) / 0.09) ** 2) for center in (0.5, 1.5, 2.5, 3.5))
        result = analyze_trace_batch(
            time,
            {manual: pulse, points: pulse * 1.1},
            {manual: AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0),
             points: AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)},
            [points], [manual], {manual},
        )
        self.assertEqual(result["summary"].iloc[0]["Reference_trace"], manual)
        self.assertEqual(result["summary"].iloc[0]["Reference_match_status"], "Matched by recording")

    def test_no_baseline_returns_zero(self):
        values = np.array([1.0, 2.0, 3.0])
        baseline = estimate_baseline(values, 0.1, AnalysisSettings(baseline_method="none"))
        np.testing.assert_array_equal(baseline, np.zeros_like(values))

    def test_constant_percentile_baseline(self):
        values = np.array([0.0, 1.0, 2.0, 100.0])
        settings = AnalysisSettings(
            baseline_method="constant_percentile",
            baseline_percentile=50.0,
        )
        np.testing.assert_allclose(estimate_baseline(values, 0.1, settings), 1.5)

    def test_prepare_signal_sorts_and_deduplicates_time(self):
        prepared = prepare_signal(
            [0.2, 0.0, 0.1, 0.1, 0.3, 0.4],
            [2.0, 0.0, 1.0, 99.0, 3.0, 4.0],
            AnalysisSettings(smoothing_s=0.0),
        )
        np.testing.assert_array_equal(prepared.t, [0.0, 0.1, 0.2, 0.3, 0.4])
        np.testing.assert_array_equal(prepared.corrected, [0.0, 1.0, 2.0, 3.0, 4.0])

    def test_prepare_signal_shifts_time_by_the_trace_frame_interval(self):
        time = np.arange(0.0, 1.0, 0.1)
        prepared = prepare_signal(
            time,
            np.linspace(0.0, 1.0, len(time)),
            AnalysisSettings(smoothing_s=0.0, time_shift_frames=-3),
        )

        self.assertAlmostEqual(prepared.t[0], -0.3)
        self.assertAlmostEqual(prepared.t[-1], 0.6)

    def test_first_detected_peak_uses_the_analysis_peak_settings(self):
        time = np.arange(0.0, 3.0, 0.1)
        signal = np.exp(-((time - 1.2) / 0.14) ** 2)
        settings = AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)

        peak_time, dt = first_detected_peak_time(time, signal, settings)
        prepared = prepare_signal(time, signal, settings)
        peaks = detect_signal_peaks(prepared.smooth, prepared.dt, settings)

        self.assertAlmostEqual(peak_time, 1.2)
        self.assertAlmostEqual(dt, 0.1)
        self.assertEqual(peaks.tolist(), [12])

    def test_align_first_peak_shifts_selected_traces_to_the_earliest_peak(self):
        page = type("PageStub", (), {})()
        time = np.arange(0.0, 3.0, 0.1)
        page.data = pd.DataFrame({
            "time_s": time,
            "A": np.exp(-((time - 1.0) / 0.14) ** 2),
            "B": np.exp(-((time - 1.4) / 0.14) ** 2),
        })
        page.time_column = "time_s"
        page.trace_processing = {
            "A": TraceProcessingSettings(),
            "B": TraceProcessingSettings(),
        }
        page.plot_cache = {"A": {}, "B": {}}
        page.selected_columns = lambda: ["A", "B"]
        page.settings = lambda: AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)
        page.settings_for_column = (
            lambda column, shared: SignalAnalysisPage.settings_for_column(page, column, shared)
        )
        page.update_trace_row_processing = lambda _column: None
        plot_requests = []
        page.schedule_plot = lambda **kwargs: plot_requests.append(kwargs)
        page.set_status = lambda _message: None

        SignalAnalysisPage.align_selected_traces_to_first_peak(page)

        self.assertEqual(page.trace_processing["A"].time_shift_frames, 0)
        self.assertEqual(page.trace_processing["B"].time_shift_frames, -4)
        self.assertNotIn("B", page.plot_cache)
        self.assertIn("A", page.plot_cache)
        self.assertEqual(plot_requests, [{"preserve_view": True}])

    def test_align_first_peak_second_click_restores_prior_offsets(self):
        page = type("PageStub", (), {})()
        time = np.arange(0.0, 3.0, 0.1)
        page.data = pd.DataFrame({
            "time_s": time,
            "A": np.exp(-((time - 1.0) / 0.14) ** 2),
            "B": np.exp(-((time - 1.4) / 0.14) ** 2),
        })
        page.time_column = "time_s"
        page.trace_processing = {
            "A": TraceProcessingSettings(time_shift_frames=2),
            "B": TraceProcessingSettings(time_shift_frames=-1),
        }
        page.plot_cache = {"A": {}, "B": {}}
        page.selected_columns = lambda: ["A", "B"]
        page.settings = lambda: AnalysisSettings(smoothing_s=0.0, prominence_pct=10.0)
        page.settings_for_column = (
            lambda column, shared: SignalAnalysisPage.settings_for_column(page, column, shared)
        )
        page.update_trace_row_processing = lambda _column: None
        page.schedule_plot = lambda **_kwargs: None
        page.set_status = lambda _message: None

        SignalAnalysisPage.align_selected_traces_to_first_peak(page)
        SignalAnalysisPage.align_selected_traces_to_first_peak(page)

        self.assertEqual(page.trace_processing["A"].time_shift_frames, 2)
        self.assertEqual(page.trace_processing["B"].time_shift_frames, -1)
        self.assertEqual(page._first_peak_alignment_restore, {})

    def test_inline_frame_arrow_shifts_only_its_trace(self):
        page = type("PageStub", (), {})()
        page.data = pd.DataFrame({"time_s": np.arange(10), "A": np.arange(10), "B": np.arange(10)})
        page.trace_processing = {
            "A": TraceProcessingSettings(time_shift_frames=0),
            "B": TraceProcessingSettings(time_shift_frames=4),
        }
        page.plot_cache = {"A": {}, "B": {}}
        page.update_trace_row_processing = lambda _column: None
        page.selected_columns = lambda: []
        page.schedule_plot = lambda **_kwargs: None
        page.set_status = lambda _message: None

        SignalAnalysisPage.shift_trace_by_frames(page, "A", -1)

        self.assertEqual(page.trace_processing["A"].time_shift_frames, -1)
        self.assertEqual(page.trace_processing["B"].time_shift_frames, 4)
        self.assertNotIn("A", page.plot_cache)
        self.assertIn("B", page.plot_cache)

    def test_point_tracking_traces_use_distance_axis(self):
        page_stub = type("PageStub", (), {"force_calibration_available": False})()
        self.assertEqual(
            SignalAnalysisPage.y_axis_label(page_stub, ["Point_1_distance_um"]),
            "Distance (µm)",
        )

    def test_point_tracking_force_traces_use_force_axis(self):
        page_stub = type("PageStub", (), {"force_calibration_available": False})()
        self.assertEqual(
            SignalAnalysisPage.y_axis_label(page_stub, ["Point_1_Force_uN"]),
            "Force (µN)",
        )

    def test_combined_point_tracking_file_gets_method_prefix(self):
        self.assertEqual(
            SignalAnalysisPage.trace_source_identity(
                Path("point_tracking_force_traces.csv")
            ),
            ("Point Tracking", ""),
        )


if __name__ == "__main__":
    unittest.main()
